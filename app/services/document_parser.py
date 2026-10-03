import re
import io
import os
from pathlib import Path
from typing import Tuple, Optional
from app.config import config


class DocumentParsingError(Exception):
    pass


class DocumentParser:
    @staticmethod
    def clean_text(text: str) -> str:
        if not text:
            return ""
        # Normalize non-breaking spaces and unusual whitespace
        text = text.replace("\u00a0", " ").replace("\u200b", "").replace("\r\n", "\n").replace("\r", "\n")
        # Remove null bytes or control characters except newlines/tabs
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
        # Normalize multiple spaces (preserve single spaces)
        text = re.sub(r"[ \t]+", " ", text)
        # Normalize 3+ consecutive newlines to 2 newlines
        text = re.sub(r"\n{3,}", "\n\n", text)
        # Strip leading and trailing whitespace
        text = text.strip()
        # Truncate if exceeds max safety limit
        if len(text) > config.MAX_EXTRACTED_CHARS:
            text = text[: config.MAX_EXTRACTED_CHARS] + "\n\n...[Content truncated for analysis safety]..."
        return text

    @classmethod
    def extract_from_pdf(cls, file_bytes: bytes) -> str:
        try:
            import fitz  # PyMuPDF
        except ImportError:
            raise DocumentParsingError("PyMuPDF (fitz) is not installed.")

        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            extracted_pages = []
            for page_num in range(len(doc)):
                page = doc[page_num]
                page_text = page.get_text("text")
                if page_text and page_text.strip():
                    extracted_pages.append(page_text.strip())
            doc.close()

            full_text = "\n\n".join(extracted_pages)
            cleaned = cls.clean_text(full_text)
            if not cleaned or len(cleaned) < 20:
                raise DocumentParsingError("Could not extract readable text from PDF. It might be scanned or image-only.")
            return cleaned
        except DocumentParsingError:
            raise
        except Exception as e:
            raise DocumentParsingError(f"Failed to parse PDF document: {str(e)}")

    @classmethod
    def extract_from_docx(cls, file_bytes: bytes) -> str:
        try:
            import docx
        except ImportError:
            raise DocumentParsingError("python-docx is not installed.")

        try:
            file_stream = io.BytesIO(file_bytes)
            doc = docx.Document(file_stream)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            full_text = "\n".join(paragraphs)
            cleaned = cls.clean_text(full_text)
            if not cleaned or len(cleaned) < 20:
                raise DocumentParsingError("Could not extract readable text from DOCX file.")
            return cleaned
        except DocumentParsingError:
            raise
        except Exception as e:
            raise DocumentParsingError(f"Failed to parse DOCX document: {str(e)}")

    @classmethod
    def extract_from_file(cls, filename: str, file_bytes: bytes) -> str:
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return cls.extract_from_pdf(file_bytes)
        elif ext in [".docx", ".doc"]:
            return cls.extract_from_docx(file_bytes)
        elif ext in [".txt", ".md", ".rtf"]:
            try:
                text = file_bytes.decode("utf-8", errors="ignore")
                cleaned = cls.clean_text(text)
                if not cleaned:
                    raise DocumentParsingError("Text file is empty.")
                return cleaned
            except Exception as e:
                raise DocumentParsingError(f"Failed to read text file: {str(e)}")
        else:
            raise DocumentParsingError(f"Unsupported file format '{ext}'. Please upload a PDF or DOCX file.")
